"""Hand-annotated validation subset for M3 (Chapter 4 Table 4.5): 600 articles stratified by language
(200 English, 200 Sinhala, 200 Tamil), of which 20% (120, 40 per language) are double-annotated by a
second person for Cohen's kappa (target >= 0.60).

  make  : python scripts/annotation.py make <news_raw folder> <out_dir>
          -> annotator_A.xlsx (600 rows), annotator_B.xlsx (the 120 overlap rows), guidelines in both.
  score : python scripts/annotation.py score <annotator_A.xlsx> <annotator_B.xlsx> <articles_scored.parquet> <out.xlsx>
          -> kappa between annotators; model accuracy / macro-F1 per language against annotator A.

Annotators see only the headline and lead text, never the model's score.
"""

from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation
from sklearn.metrics import accuracy_score, cohen_kappa_score, confusion_matrix, f1_score

SEED = 42
PER_LANG, OVERLAP_PER_LANG = 200, 40
GUIDE = [
    "Label the article's likely effect on the share price of the company it is about (or on the market",
    "as a whole if no single company is named), as a careful investor reading it on the day would see it.",
    "",
    "positive = good news for the price (higher profit, dividend, new contract, rate cut, IMF tranche ...)",
    "negative = bad news for the price (loss, downgrade, tax rise, crisis, fraud, strike ...)",
    "neutral  = no clear price direction (routine notice, appointment, mixed or purely factual report)",
    "",
    "Decide from the headline and lead shown. Do not look up later prices. If unsure, choose neutral.",
    "Work alone; do not discuss items with the other annotator until both files are finished.",
]


def load(src: Path) -> pd.DataFrame:
    """Scraped articles (news_scraper.py output folder); only items with fetched text are sampled."""
    # article files only (not headline lists, not the *_refetched.csv side files whose bodies are
    # already merged into hiru_<lang>_articles.csv)
    files = sorted(f for f in src.rglob("*.csv") if f.name.startswith("adaderana") or f.name.endswith("_articles.csv"))
    frames = [pd.read_csv(f, dtype=str) for f in files]
    d = pd.concat(frames, ignore_index=True).drop_duplicates("url")
    d["lang_group"] = d["language_hint"].map({"en": "en", "english": "en", "sinhala": "si", "tamil": "ta"})
    return d


def make(src: str, out: str):
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    d = load(Path(src))
    if "body" not in d:
        d["body"] = ""
    d["lead"] = d["body"].fillna("").str[:400]
    d = d[d["lead"].str.len() > 40]
    picks = []
    for lang in ("en", "si", "ta"):
        g = d[d["lang_group"] == lang]
        if len(g) < PER_LANG:
            print(f"warning: only {len(g)} {lang} articles available")
        picks.append(g.sample(min(PER_LANG, len(g)), random_state=SEED).assign(lang_group=lang))
    s = pd.concat(picks).sample(frac=1, random_state=SEED).reset_index(drop=True)
    s.insert(0, "item_id", [f"A{i:04d}" for i in range(1, len(s) + 1)])
    s["label"] = ""
    cols = ["item_id", "lang_group", "source", "published_at", "title", "lead", "url", "label"]
    s = s.reindex(columns=cols)
    overlap = pd.concat(s[s["lang_group"] == l].head(OVERLAP_PER_LANG) for l in ("en", "si", "ta"))
    for name, frame in (("annotator_A.xlsx", s), ("annotator_B.xlsx", overlap)):
        p = out / name
        with pd.ExcelWriter(p) as xw:
            frame.to_excel(xw, sheet_name="items", index=False)
            pd.DataFrame({"Guidelines": GUIDE}).to_excel(xw, sheet_name="guidelines", index=False)
        wb = load_workbook(p)
        ws = wb["items"]
        dv = DataValidation(type="list", formula1='"positive,neutral,negative"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"H2:H{len(frame) + 1}")
        for col, width in zip("ABCDEFGH", (9, 6, 16, 18, 60, 80, 30, 12)):
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows(min_row=2):
            for c in row[4:6]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws["H1"].font = Font(bold=True)
        ws.freeze_panes = "A2"
        wb.save(p)
        print(p, len(frame))


def score(a: str, b: str, scored: str, out: str):
    A = pd.read_excel(a, sheet_name="items")
    B = pd.read_excel(b, sheet_name="items")
    m = A.merge(B[["item_id", "label"]], on="item_id", suffixes=("_a", "_b")).dropna(subset=["label_a", "label_b"])
    rows = [{"scope": "all", "n": len(m), "cohen_kappa": cohen_kappa_score(m["label_a"], m["label_b"]),
             "raw_agreement": float((m["label_a"] == m["label_b"]).mean())}]
    for l, g in m.groupby("lang_group"):
        rows.append({"scope": l, "n": len(g), "cohen_kappa": cohen_kappa_score(g["label_a"], g["label_b"]),
                     "raw_agreement": float((g["label_a"] == g["label_b"]).mean())})
    kappa = pd.DataFrame(rows)
    S = pd.read_parquet(scored)[["url", "lang", "model", "p_neg", "p_neu", "p_pos", "fallback_used"]]
    S["pred"] = np.array(["negative", "neutral", "positive"])[S[["p_neg", "p_neu", "p_pos"]].values.argmax(1)]
    g = A.dropna(subset=["label"]).merge(S, on="url", how="inner")
    res = []
    for l, x in [("all", g)] + list(g.groupby("lang_group")):
        res.append({"language": l, "n": len(x), "accuracy": accuracy_score(x["label"], x["pred"]),
                    "macro_f1": f1_score(x["label"], x["pred"], average="macro", zero_division=0),
                    "fallback_share": float(x["fallback_used"].mean())})
    cm = pd.DataFrame(confusion_matrix(g["label"], g["pred"], labels=["negative", "neutral", "positive"]),
                      index=["gold negative", "gold neutral", "gold positive"], columns=["pred negative", "pred neutral", "pred positive"])
    with pd.ExcelWriter(out) as xw:
        kappa.to_excel(xw, sheet_name="Inter-annotator kappa", index=False)
        pd.DataFrame(res).to_excel(xw, sheet_name="Model vs gold", index=False)
        cm.to_excel(xw, sheet_name="Confusion matrix")
    print(kappa.round(3).to_string(index=False))
    print(pd.DataFrame(res).round(3).to_string(index=False))


if __name__ == "__main__":
    if sys.argv[1] == "make":
        make(sys.argv[2], sys.argv[3])
    else:
        score(*sys.argv[2:6])
