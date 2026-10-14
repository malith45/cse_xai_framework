"""Tests for src/loaders/public_float.py: unit detection, the range assertion, quarter parsing,
and sanity checks on the real files."""

import pandas as pd
import pytest
from openpyxl import Workbook

from loaders import (
    Convention,
    DuplicateRowsError,
    FloatOutOfRange,
    PublicFloatError,
    QuarterParseError,
    audit_duplicates,
    detect_numeric_convention,
    load_public_float,
    load_public_float_file,
    normalise_float_values,
    parse_quarter_code,
    quarter_end_date,
)

from .conftest import PUBLIC_HOLDING, requires_data


def _workbook(tmp_path, sheet, rows, title=None):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    if title:
        ws.append([title])
        ws.append([])
    ws.append(["QUARTER", "COMPANY_NAME", "SECURITY", "PUBLIC_FLOAT"])
    for row in rows:
        ws.append(row)
    path = tmp_path / "Public Holding 2099.xlsx"
    wb.save(path)
    return path


# ------------------------------------------------------------------ quarter parsing
class TestQuarterParsing:
    @pytest.mark.parametrize(
        "code,expected",
        [
            ("2012-4", (2012, 4)),
            ("2018-1", (2018, 1)),
            ("2025-01", (2025, 1)),
            ("2025-04", (2025, 4)),
            ("2016 - 1", (2016, 1)),
            (" 2023-03 ", (2023, 3)),
        ],
    )
    def test_all_formats(self, code, expected):
        assert parse_quarter_code(code) == expected

    @pytest.mark.parametrize("bad", ["2025-05", "2025-13", "2025", "Q1-2025", "", None, float("nan"), "2025-0"])
    def test_rejects(self, bad):
        with pytest.raises(ValueError):
            parse_quarter_code(bad)

    def test_quarter_end_dates(self):
        assert quarter_end_date(2024, 1) == pd.Timestamp("2024-03-31")
        assert quarter_end_date(2024, 2) == pd.Timestamp("2024-06-30")
        assert quarter_end_date(2024, 3) == pd.Timestamp("2024-09-30")
        assert quarter_end_date(2024, 4) == pd.Timestamp("2024-12-31")
        with pytest.raises(ValueError):
            quarter_end_date(2024, 5)


# ------------------------------------------------------------------ unit detection (synthetic)
class TestUnitDetection:
    def test_fraction_convention(self):
        s = pd.Series([0.1024, 0.5, 1.0, 0, 0.0034])
        out, conv = normalise_float_values(s, file="f", sheet="s")
        assert conv is Convention.FRACTION
        assert out.tolist() == pytest.approx([10.24, 50.0, 100.0, 0.0, 0.34])

    def test_percent_string_convention(self):
        s = pd.Series(["2.84%", "100%", "0%", " 27.04% "])
        out, conv = normalise_float_values(s, file="f", sheet="s")
        assert conv is Convention.PERCENT_STRING
        assert out.tolist() == pytest.approx([2.84, 100.0, 0.0, 27.04])

    def test_percent_points_numeric_convention(self):
        s = pd.Series([2.84, 100, 0, 27.04])
        out, conv = normalise_float_values(s, file="f", sheet="s")
        assert conv is Convention.PERCENT_POINTS
        assert out.tolist() == pytest.approx([2.84, 100.0, 0.0, 27.04])

    def test_mixed_sheet_uses_per_value_rule(self):
        # 2020 Q2 shape: fractions (including int 0 and 1) alongside percent strings
        s = pd.Series([0.352, "28.62%", 1, 0, "0.5%", 0.2])
        out, conv = normalise_float_values(s, file="f", sheet="s")
        assert conv is Convention.MIXED
        assert out.tolist() == pytest.approx([35.2, 28.62, 100.0, 0.0, 0.5, 20.0])

    def test_detect_numeric_rule_is_distribution_based(self):
        assert detect_numeric_convention([0.1, 0.9, 1.0]) is Convention.FRACTION
        assert detect_numeric_convention([0.1, 1.01]) is Convention.PERCENT_POINTS
        assert detect_numeric_convention(["10%", "20%"]) is Convention.EMPTY
        assert detect_numeric_convention([]) is Convention.EMPTY

    def test_nan_and_bool_are_not_numbers(self):
        with pytest.raises(FloatOutOfRange):
            normalise_float_values(pd.Series([0.2, float("nan"), 0.3]), file="f", sheet="s")
        with pytest.raises(FloatOutOfRange):
            normalise_float_values(pd.Series([0.2, True, 0.3]), file="f", sheet="s")


# ------------------------------------------------------------------ range assertion
class TestRangeAssertion:
    def test_above_100_names_file_sheet_row_and_security(self):
        s = pd.Series([25.0, 101.5, 30.0], index=[10, 11, 12])
        labels = pd.Series(["A-N-0000", "B-N-0000", "C-N-0000"], index=s.index)
        with pytest.raises(FloatOutOfRange) as info:
            normalise_float_values(s, file="Public Holding 2099.xlsx", sheet="1-2099", row_labels=labels)
        msg = str(info.value)
        assert "Public Holding 2099.xlsx" in msg and "1-2099" in msg
        assert "row 11" in msg and "B-N-0000" in msg and "101.5" in msg

    def test_negative_fraction_halts(self):
        with pytest.raises(FloatOutOfRange) as info:
            normalise_float_values(pd.Series([0.3, -0.01, 0.25]), file="f.xlsx", sheet="s")
        assert "row 1" in str(info.value)

    def test_percent_string_out_of_range_halts(self):
        with pytest.raises(FloatOutOfRange):
            normalise_float_values(pd.Series(["25%", "250%"]), file="f.xlsx", sheet="s")

    def test_unreadable_text_halts(self):
        with pytest.raises(FloatOutOfRange) as info:
            normalise_float_values(pd.Series([0.3, "n/a"]), file="f.xlsx", sheet="s")
        assert "n/a" in str(info.value)

    def test_nothing_is_clamped(self):
        with pytest.raises(FloatOutOfRange):
            normalise_float_values(pd.Series([25.0, 100.0001]), file="f", sheet="s")

    def test_implausible_median_halts(self):
        # one value above 1 forces the points rule; the remaining tiny values then give a median far
        # below any real sheet, which must be reported rather than returned
        with pytest.raises(FloatOutOfRange) as info:
            normalise_float_values(pd.Series([0.2, 0.3, 0.25, 1.5]), file="f", sheet="s")
        assert "median" in str(info.value)

    def test_end_to_end_bad_value_halts_with_location(self, tmp_path):
        path = _workbook(
            tmp_path, "1-2099",
            [["2099-01", "GOOD CO PLC", "GOOD.N0000", "25.5%"], ["2099-01", "BAD CO PLC", "BADC.N0000", "125.5%"]],
            title="Public Holding 2099Q1 (DATA as per latest CSE records)",
        )
        with pytest.raises(FloatOutOfRange) as info:
            load_public_float_file(path)
        msg = str(info.value)
        assert "Public Holding 2099.xlsx" in msg and "1-2099" in msg and "BADC-N-0000" in msg

    def test_end_to_end_quarter_mismatch_halts(self, tmp_path):
        path = _workbook(tmp_path, "2-2099", [["2099-01", "GOOD CO PLC", "GOOD.N0000", "25.5%"]])
        with pytest.raises(QuarterParseError):
            load_public_float_file(path)

    def test_end_to_end_bare_ticker_halts(self, tmp_path):
        path = _workbook(tmp_path, "1-2099", [["2099-01", "GOOD CO PLC", "GOOD", "25.5%"]])
        with pytest.raises(PublicFloatError) as info:
            load_public_float_file(path)
        assert "share class" in str(info.value)

    def test_end_to_end_good_file_round_trips(self, tmp_path):
        path = _workbook(
            tmp_path, "3-2099",
            [["2099-3", "GOOD CO PLC", "GOOD.N0000", 0.255], ["2099-3", "GOOD CO PLC", "good.x0000", 0.10]],
        )
        df, reports = load_public_float_file(path)
        assert df["security_id"].tolist() == ["GOOD-N-0000", "GOOD-X-0000"]
        assert df["float_pct"].tolist() == pytest.approx([25.5, 10.0])
        assert df["quarter_end"].tolist() == [pd.Timestamp("2099-09-30")] * 2
        assert reports[0].convention is Convention.FRACTION


# ------------------------------------------------------------------ real files
@pytest.fixture(scope="module")
def public_float():
    if not PUBLIC_HOLDING.is_dir():
        pytest.skip("CSE Data Library not found")
    return load_public_float(PUBLIC_HOLDING, with_reports=True, duplicate_policy="prefer_varying_value")


def _multi_sheet_workbook(tmp_path, sheets):
    """sheets: {sheet_name: [rows]} -> one workbook in the public-holding layout."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(["QUARTER", "COMPANY_NAME", "SECURITY", "PUBLIC_FLOAT"])
        for row in rows:
            ws.append(row)
    path = tmp_path / "Public Holding 2099.xlsx"
    wb.save(path)
    return path


class TestPreferVaryingValuePolicy:
    """Synthetic cases. A stale duplicate is the one that does not move across quarters; the ID
    spelling plays no part in the decision."""

    def test_frozen_row_dropped_live_row_kept_regardless_of_case(self, tmp_path):
        # the FROZEN row is the upper-case one here, so a casing rule would pick wrongly
        filler = lambda q: [[f"2099-0{q}", "A PLC", "AAAA.N0000", "20%"], [f"2099-0{q}", "B PLC", "BBBB.N0000", "25%"],
                            [f"2099-0{q}", "C PLC", "CCCC.N0000", "30%"]]
        path = _multi_sheet_workbook(tmp_path, {
            "1-2099": [["2099-01", "DUP PLC", "DUPL.N0000", "70.29%"], ["2099-01", "DUP PLC", "dupl.n0000", "70.5%"]] + filler(1),
            "2-2099": [["2099-02", "DUP PLC", "DUPL.N0000", "70.29%"], ["2099-02", "DUP PLC", "dupl.n0000", "70.64%"]] + filler(2),
            "3-2099": [["2099-03", "DUP PLC", "DUPL.N0000", "70.29%"], ["2099-03", "DUP PLC", "Dupl.N0000", "72.58%"]] + filler(3),
        })
        df, reports = load_public_float_file(path, duplicate_policy="prefer_varying_value")
        dup = df[df["security_id"] == "DUPL-N-0000"].sort_values("quarter")
        assert dup["float_pct"].tolist() == pytest.approx([70.5, 70.64, 72.58])
        assert len(df) == 12
        for rep in reports:
            assert rep.rows == 4
            assert "frozen at 70.29%" in rep.notes and "'DUPL.N0000'" in rep.notes

    def test_both_vary_halts_and_prints_both_rows(self, tmp_path):
        path = _multi_sheet_workbook(tmp_path, {
            "1-2099": [["2099-01", "DUP PLC", "DUPL.N0000", "30%"], ["2099-01", "DUP PLC", "dupl.n0000", "40%"]],
            "2-2099": [["2099-02", "DUP PLC", "DUPL.N0000", "31%"], ["2099-02", "DUP PLC", "dupl.n0000", "41%"]],
        })
        with pytest.raises(DuplicateRowsError) as info:
            load_public_float_file(path, duplicate_policy="prefer_varying_value")
        msg = str(info.value)
        assert "both vary" in msg
        assert "'DUPL.N0000' -> '30%'" in msg and "'dupl.n0000' -> '40%'" in msg and "'dupl.n0000' -> '41%'" in msg
        assert "1-2099" in msg and "2-2099" in msg

    def test_both_frozen_halts_and_prints_both_rows(self, tmp_path):
        path = _multi_sheet_workbook(tmp_path, {
            "1-2099": [["2099-01", "DUP PLC", "DUPL.N0000", "30%"], ["2099-01", "DUP PLC", "dupl.n0000", "40%"]],
            "2-2099": [["2099-02", "DUP PLC", "DUPL.N0000", "30%"], ["2099-02", "DUP PLC", "dupl.n0000", "40%"]],
        })
        with pytest.raises(DuplicateRowsError) as info:
            load_public_float_file(path, duplicate_policy="prefer_varying_value")
        msg = str(info.value)
        assert "both frozen" in msg and "30.0" in msg and "40.0" in msg
        assert "'DUPL.N0000' -> '30%'" in msg and "'dupl.n0000' -> '40%'" in msg

    def test_single_duplicated_quarter_halts(self, tmp_path):
        path = _multi_sheet_workbook(tmp_path, {
            "1-2099": [["2099-01", "DUP PLC", "DUPL.N0000", "30%"], ["2099-01", "DUP PLC", "dupl.n0000", "40%"]],
            "2-2099": [["2099-02", "DUP PLC", "DUPL.N0000", "31%"]],
        })
        with pytest.raises(DuplicateRowsError) as info:
            load_public_float_file(path, duplicate_policy="prefer_varying_value")
        assert "single quarter" in str(info.value)

    def test_default_policy_is_halt(self, tmp_path):
        filler = lambda q: [[f"2099-0{q}", "A PLC", "AAAA.N0000", "20%"], [f"2099-0{q}", "B PLC", "BBBB.N0000", "25%"],
                            [f"2099-0{q}", "C PLC", "CCCC.N0000", "30%"]]
        path = _multi_sheet_workbook(tmp_path, {
            "1-2099": [["2099-01", "DUP PLC", "DUPL.N0000", "70.29%"], ["2099-01", "DUP PLC", "dupl.n0000", "70.5%"]] + filler(1),
            "2-2099": [["2099-02", "DUP PLC", "DUPL.N0000", "70.29%"], ["2099-02", "DUP PLC", "dupl.n0000", "70.64%"]] + filler(2),
        })
        with pytest.raises(DuplicateRowsError) as info:
            load_public_float_file(path)
        assert "policy 'error'" in str(info.value) and "'DUPL.N0000' -> '70.29%'" in str(info.value)
        with pytest.raises(ValueError):
            load_public_float_file(path, duplicate_policy="prefer_uppercase_id")

    def test_no_duplicates_is_untouched(self, tmp_path):
        path = _multi_sheet_workbook(tmp_path, {
            "1-2099": [["2099-01", "A PLC", "AAAA.N0000", "30%"]],
            "2-2099": [["2099-02", "A PLC", "AAAA.N0000", "31%"]],
        })
        df, reports = load_public_float_file(path, duplicate_policy="prefer_varying_value")
        assert len(df) == 2 and all(r.notes == "" for r in reports)


@requires_data
class TestRealDuplicateRows:
    """2024 Q3 - 2025 Q4 list VIDULLANKA's non-voting share twice: one row frozen at 70.29% and one
    that moves every quarter. Both spellings canonicalise to VLL-X-0000."""

    def test_default_policy_halts_and_names_both_rows(self):
        with pytest.raises(DuplicateRowsError) as info:
            load_public_float(PUBLIC_HOLDING)
        msg = str(info.value)
        assert "VLL-X-0000" in msg and "Public Holding 2024.xlsx" in msg and "3-2024" in msg
        assert "'VLL.x0000' -> '70.29%'" in msg and "'VLL.X0000' -> '70.5%'" in msg

    def test_varying_row_kept_frozen_row_dropped_and_recorded(self, public_float):
        df, reports = public_float
        vll = df[df["security_id"] == "VLL-X-0000"].set_index(["year", "quarter"])["float_pct"]
        assert vll[(2024, 3)] == pytest.approx(70.5)
        assert vll[(2024, 4)] == pytest.approx(70.64)
        assert vll[(2025, 4)] == pytest.approx(72.85)
        noted = [r for r in reports if "duplicates:" in r.notes]
        assert [(r.file, r.sheet) for r in noted] == [
            ("Public Holding 2024.xlsx", "3-2024"), ("Public Holding 2024.xlsx", "4-2024"),
            ("Public Holding 2025.xlsx", "1-2025"), ("Public Holding 2025.xlsx", "2-2025"),
            ("Public Holding 2025.xlsx", "3-2025"), ("Public Holding 2025.xlsx", "4-2025"),
        ]
        assert all("frozen at 70.29% across 6 quarters" in r.notes for r in noted)

    def test_audit_lists_every_candidate_row(self):
        audit = audit_duplicates(PUBLIC_HOLDING)
        assert len(audit) == 12
        assert set(audit["security_id"]) == {"VLL-X-0000"}
        assert set(audit["raw_id"]) == {"VLL.x0000", "VLL.X0000"}
        frozen = audit[audit["raw_float"] == "70.29%"]
        assert len(frozen) == 6


@requires_data
class TestRealFiles:
    def test_schema_and_coverage(self, public_float):
        df, reports = public_float
        assert list(df.columns) == ["security_id", "company_name", "year", "quarter", "quarter_end", "float_pct"]
        assert df["year"].min() == 2012 and df["year"].max() == 2025
        assert set(df["quarter"]) == {1, 2, 3, 4}
        assert len(reports) == 56
        assert len(df) > 17000
        assert not df.duplicated(["security_id", "year", "quarter"]).any()
        assert df["float_pct"].between(0, 100).all()
        assert (df["quarter_end"].dt.month == df["quarter"] * 3).all()

    def test_2012_fraction_and_2018_percent_detected_from_values(self):
        _, r2012 = load_public_float_file(PUBLIC_HOLDING / "Public Holding 2012.xlsx")
        _, r2018 = load_public_float_file(PUBLIC_HOLDING / "Public Holding 2018.xlsx")
        assert [r.convention for r in r2012] == [Convention.FRACTION] * 4
        assert [r.sheet for r in r2012] == ["2012 Q1", "2012 Q2", "2012 Q3", "2012 Q4"]  # reverse order fixed
        assert {r.sheet: r.convention for r in r2018} == {
            "2018_Q1": Convention.PERCENT_STRING,
            "2018_Q2": Convention.PERCENT_STRING,
            "2018_Q3": Convention.FRACTION,
            "2018_Q4": Convention.PERCENT_STRING,
        }

    def test_known_values_normalised_identically_across_conventions(self, public_float):
        df, _ = public_float
        by = df.set_index(["security_id", "year", "quarter"])["float_pct"]
        assert by[("ACAP-N-0000", 2012, 4)] == pytest.approx(10.24)   # stored 0.1024
        assert by[("MFL-N-0000", 2019, 4)] == pytest.approx(35.2)     # "35.2%"
        assert by[("MFL-N-0000", 2020, 2)] == pytest.approx(35.2)     # 0.352 inside the mixed sheet
        assert by[("MFL-N-0000", 2020, 3)] == pytest.approx(35.2)     # 0.352 fraction sheet
        assert by[("BLUE-X-0000", 2020, 2)] == pytest.approx(100.0)   # integer 1 inside the mixed sheet
        assert by[("AAF-N-0000", 2025, 1)] == pytest.approx(27.04)    # "27.04%"

    def test_2020_q1_covid_note_gives_empty_quarter(self, public_float):
        df, reports = public_float
        r = next(r for r in reports if r.file == "Public Holding 2020.xlsx" and r.sheet == "2020_Q1")
        assert r.convention is Convention.EMPTY and r.rows == 0
        assert "No records" in r.notes
        assert ((df["year"] == 2020) & (df["quarter"] == 1)).sum() == 0

    def test_2020_q2_mixed_sheet(self, public_float):
        _, reports = public_float
        r = next(r for r in reports if r.sheet == "2020_Q2")
        assert r.convention is Convention.MIXED and r.rows == 341

    def test_ids_are_canonical_and_classes_preserved(self, public_float):
        df, _ = public_float
        assert df["security_id"].str.fullmatch(r"[A-Z0-9]{2,5}-[A-Z]-\d{4}").all()
        assert "VLL-X-0000" in set(df["security_id"])  # from "VLL.x0000"
        assert {"COMB-N-0000", "COMB-X-0000"} <= set(df["security_id"])

    def test_float_does_not_jump_implausibly_between_consecutive_quarters(self, public_float):
        # A unit misdetection in any sheet would multiply or divide every float by 100 at that
        # quarter boundary. Check consecutive-quarter changes per security, overall and at each
        # convention boundary.
        df, _ = public_float
        d = df.sort_values(["security_id", "year", "quarter"]).copy()
        d["t"] = d["year"] * 4 + d["quarter"]
        prev = d.groupby("security_id").shift(1)
        consecutive = (d["t"] - prev["t"]) == 1
        d["change"] = (d["float_pct"] - prev["float_pct"]).abs()
        c = d.loc[consecutive]
        assert len(c) > 15000
        assert c["change"].median() < 0.5
        assert c["change"].quantile(0.95) < 10
        assert c["change"].quantile(0.99) < 40
        boundaries = {(2016, 4), (2018, 3), (2018, 4), (2020, 3), (2020, 4), (2022, 3), (2022, 4)}
        for year, quarter in sorted(boundaries):
            med = c.loc[(c["year"] == year) & (c["quarter"] == quarter), "change"].median()
            assert med < 0.5, f"median change entering {year}Q{quarter} is {med}"
